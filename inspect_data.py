"""Inspect online_retail_II.xlsx workbook in read-only mode."""
import openpyxl
from pathlib import Path
import json

WORKBOOK = Path("data/raw/online_retail_II.xlsx")

wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
print(f"Sheet names: {wb.sheetnames}")

info = {}
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== Sheet: {name} ===")
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    print(f"Columns: {header}")
    print(f"Dimensions: {ws.dimensions}, max_row={ws.max_row}")
    sample = []
    for i, row in enumerate(rows):
        if i < 3:
            sample.append(row)
        else:
            if i == ws.max_row - 2:
                pass
    print(f"First 3 rows:")
    for r in sample:
        print(f"  {r}")
    info[name] = {"columns": list(header), "max_row": ws.max_row, "sample": [list(r) for r in sample]}
wb.close()

with open("/tmp/workbook_info.json", "w") as f:
    json.dump(info, f, default=str, indent=2)
print("\nSaved /tmp/workbook_info.json")
